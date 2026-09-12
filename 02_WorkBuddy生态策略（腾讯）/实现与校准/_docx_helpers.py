# -*- coding: utf-8 -*-
"""docx 生成公共助手（遵循 docx-generation 技能设计规范：方案文档 / blue 主色）。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE_800 = RGBColor(0x0C, 0x44, 0x7C)
BLUE_600 = RGBColor(0x18, 0x5F, 0xA5)
BLUE_400 = RGBColor(0x37, 0x8A, 0xDD)
GRAY_900 = RGBColor(0x2C, 0x2C, 0x2A)
GRAY_600 = RGBColor(0x5F, 0x5E, 0x5A)
PURPLE_600 = RGBColor(0x53, 0x4A, 0xB7)
GREEN_600 = RGBColor(0x63, 0x99, 0x22)
AMBER_600 = RGBColor(0xBA, 0x75, 0x17)
RED_600 = RGBColor(0xE2, 0x4B, 0x4A)
HEI = "黑体"
SONG = "宋体"


def make_doc():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2.54); sec.bottom_margin = Cm(2.54)
        sec.left_margin = Cm(3.17); sec.right_margin = Cm(3.17)
    normal = doc.styles["Normal"]
    normal.font.name = SONG; normal.font.size = Pt(12); normal.font.color.rgb = GRAY_900
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
    return doc


def set_run(run, font=SONG, size=12, color=GRAY_900, bold=False, italic=False):
    run.font.name = font; run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font); return run


def add_runs(p, text, size=12, color=GRAY_900, bold=False, italic=False):
    """支持 **粗体** 内联标记：把文本按 ** 切分，奇数段加粗。"""
    parts = text.split("**")
    for i, chunk in enumerate(parts):
        if chunk == "":
            continue
        set_run(p.add_run(chunk), size=size, color=color, bold=(bold or i % 2 == 1), italic=italic)
    return p


def para(doc, text, size=12, color=GRAY_900, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         indent=True, before=6, after=6, italic=False):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.alignment = align; pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.5
    if indent: pf.first_line_indent = Pt(24)
    add_runs(p, text, size=size, color=color, bold=bold, italic=italic); return p


def h1(doc, text):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER; pf.space_before = Pt(16); pf.space_after = Pt(12)
    pf.line_spacing = 1.5; pf.keep_with_next = True
    set_run(p.add_run(text), font=HEI, size=18, color=BLUE_800, bold=True); return p


def h2(doc, text):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(12); pf.space_after = Pt(6); pf.line_spacing = 1.5; pf.keep_with_next = True
    set_run(p.add_run(text), font=HEI, size=15, color=BLUE_600, bold=True); return p


def h3(doc, text):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(8); pf.space_after = Pt(4); pf.line_spacing = 1.5; pf.keep_with_next = True
    set_run(p.add_run(text), font=HEI, size=14, color=BLUE_400, bold=True); return p


def bullet(doc, text, level=0, bold_lead=None):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.left_indent = Pt(24 + 18 * level); pf.space_before = Pt(2); pf.space_after = Pt(2); pf.line_spacing = 1.5
    lead = "• " if level == 0 else "– "
    if bold_lead:
        set_run(p.add_run(lead + bold_lead + "："), color=BLUE_600, bold=True)
        add_runs(p, text)
    else:
        add_runs(p, lead + text)
    return p


def callout(doc, text, color=PURPLE_600):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(8); pf.space_after = Pt(8); pf.left_indent = Pt(12); pf.right_indent = Pt(12); pf.line_spacing = 1.5
    set_run(p.add_run(text), color=color, bold=True)
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr"); left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "12"); left.set(qn("w:space"), "8"); left.set(qn("w:color"), "534AB7")
    pbdr.append(left); pPr.append(pbdr); return p


def add_table(doc, headers, rows, col_widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""; p = hdr[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(htext), font=HEI, size=10.5, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
        shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "185FA5"); hdr[i]._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""; p = cells[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(str(val)), size=10.5)
    if col_widths:
        for i, wd in enumerate(col_widths):
            for row in t.rows: row.cells[i].width = Cm(wd)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4); return t


def add_page_number(doc):
    for sec in doc.sections:
        footer = sec.footer; p = footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(); f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1); run._r.append(it); run._r.append(f2); set_run(run, size=10.5, color=GRAY_600)


def cover(doc, title, subtitle, tagline, info_rows):
    tp = doc.add_paragraph(); tp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; tp.paragraph_format.space_before = Pt(84)
    set_run(tp.add_run(title), font=HEI, size=26, color=BLUE_800, bold=True)
    sp = doc.add_paragraph(); sp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp.paragraph_format.space_after = Pt(40)
    set_run(sp.add_run(subtitle), size=13, color=BLUE_600, bold=True)
    if tagline:
        tg = doc.add_paragraph(); tg.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; tg.paragraph_format.space_after = Pt(80)
        set_run(tg.add_run(tagline), size=12, color=GRAY_600, italic=True)
    add_table(doc, ["项目", "内容"], info_rows, col_widths=[3.2, 11.5])
    doc.add_paragraph()

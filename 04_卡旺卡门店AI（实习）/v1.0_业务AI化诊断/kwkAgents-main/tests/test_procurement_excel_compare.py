import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from backend.procurement.excel_compare import (
    WorkbookValidationError,
    compare_quote_workbook,
    create_standard_snapshot,
    write_comparison_report,
)


class ProcurementExcelCompareTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.standard_path = self.root / "standard.xlsx"
        self._write_standard(self.standard_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    @staticmethod
    def _write_standard(path: Path):
        workbook = Workbook()
        ws = workbook.active
        ws.title = "装饰直接费"
        ws.append(["序号", "项目名称", "项目特征描述", "单位", "数量", "价格", "合计"])
        ws.append([1, "墙面施工", "标准施工", "㎡", None, 100, "=E2*F2"])
        ws.append([2, "待定材料", "甲方尚未定价", "件", None, None, "=E3*F3"])
        ws.append([3, "精度测试", "按人民币分比较", "项", None, "90.004", "=E4*F4"])
        ws.append([4, "小计", None, None, None, None, "=SUM(G2:G4)"])
        workbook.save(path)

    def _quote_copy(self, name: str = "quote.xlsx") -> tuple[Path, object]:
        quote_path = self.root / name
        workbook = load_workbook(self.standard_path)
        return quote_path, workbook

    def test_compares_price_changes_pending_baseline_and_appended_items(self):
        snapshot = create_standard_snapshot(self.standard_path)
        quote_path, workbook = self._quote_copy()
        ws = workbook["装饰直接费"]
        ws["F2"] = 110
        ws["F3"] = 50
        ws["F4"] = "90.005"
        ws.append([5, "新增灯具", "供应商追加", "套", 1, 200, "=E6*F6"])
        workbook.save(quote_path)

        result = compare_quote_workbook(quote_path, snapshot)

        self.assertEqual(result["summary"]["increase"], 2)
        self.assertEqual(result["summary"]["pending_baseline"], 1)
        self.assertEqual(result["summary"]["new_item"], 1)
        by_name = {item["name"]: item for item in result["items"]}
        self.assertEqual(by_name["墙面施工"]["difference"], "10.00")
        self.assertEqual(by_name["墙面施工"]["percentage"], "10.00")
        self.assertIsNone(by_name["待定材料"]["percentage"])
        self.assertEqual(by_name["精度测试"]["difference"], "0.01")

    def test_rejects_modified_existing_item_identity(self):
        snapshot = create_standard_snapshot(self.standard_path)
        quote_path, workbook = self._quote_copy("invalid.xlsx")
        workbook["装饰直接费"]["B2"] = "被修改的项目"
        workbook.save(quote_path)

        with self.assertRaisesRegex(WorkbookValidationError, "项目结构被修改"):
            compare_quote_workbook(quote_path, snapshot)

    def test_report_keeps_unchanged_items(self):
        snapshot = create_standard_snapshot(self.standard_path)
        result = compare_quote_workbook(self.standard_path, snapshot)
        output = self.root / "report.xlsx"

        write_comparison_report(result, output)

        workbook = load_workbook(output, read_only=True)
        ws = workbook["报价对比报告"]
        self.assertEqual(ws.max_row, len(result["items"]) + 1)
        statuses = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
        self.assertIn("未变化", statuses)
        workbook.close()


if __name__ == "__main__":
    unittest.main()

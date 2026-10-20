import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.api.routes.procurement import router
from backend.db.models import User
from backend.infra.auth import get_current_user, get_db, require_admin
from backend.infra.database import Base


class ProcurementRouteTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(engine)
        self.Session = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
        with self.Session() as db:
            self.admin = User(username="admin", password_hash="test", role="admin")
            db.add(self.admin)
            db.commit()
            db.refresh(self.admin)

        app = FastAPI()
        app.include_router(router)

        def override_db():
            db = self.Session()
            try:
                yield db
            finally:
                db.close()

        def override_user():
            with self.Session() as db:
                return db.query(User).filter(User.id == self.admin.id).first()

        app.dependency_overrides[get_db] = override_db
        app.dependency_overrides[get_current_user] = override_user
        app.dependency_overrides[require_admin] = override_user
        self.client = TestClient(app)
        self.storage_patch = patch(
            "backend.api.routes.procurement.create_storage_directory",
            side_effect=self._storage_directory,
        )
        self.project_root_patch = patch("backend.api.routes.procurement.PROJECT_ROOT", self.root)
        self.storage_patch.start()
        self.project_root_patch.start()

    def tearDown(self):
        self.project_root_patch.stop()
        self.storage_patch.stop()
        self.temp_dir.cleanup()

    def _storage_directory(self, category: str) -> Path:
        directory = self.root / "data" / "procurement" / category / "record"
        suffix = 1
        candidate = directory
        while candidate.exists():
            suffix += 1
            candidate = directory.with_name(f"record-{suffix}")
        candidate.mkdir(parents=True)
        return candidate

    @staticmethod
    def _workbook_bytes(price: int) -> bytes:
        workbook = Workbook()
        ws = workbook.active
        ws.title = "装饰直接费"
        ws.append(["序号", "项目名称", "项目特征描述", "单位", "数量", "价格", "合计"])
        ws.append([1, "墙面施工", "标准施工", "㎡", 10, price, "=E2*F2"])
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_standard_upload_comparison_history_and_report_download(self):
        standard_response = self.client.post(
            "/procurement/standards",
            data={"version_name": "2026-Q3"},
            files={"file": ("standard.xlsx", self._workbook_bytes(100), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(standard_response.status_code, 200, standard_response.text)
        standard_id = standard_response.json()["id"]
        self.assertEqual(standard_response.json()["item_count"], 1)

        comparison_response = self.client.post(
            "/procurement/comparisons",
            data={"standard_version_id": str(standard_id)},
            files={"file": ("quote.xlsx", self._workbook_bytes(115), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(comparison_response.status_code, 200, comparison_response.text)
        payload = comparison_response.json()
        self.assertEqual(payload["summary"]["increase"], 1)
        self.assertEqual(payload["items"][0]["difference"], "15.00")

        history_response = self.client.get("/procurement/comparisons")
        self.assertEqual(history_response.status_code, 200)
        self.assertEqual(len(history_response.json()), 1)

        report_response = self.client.get(f"/procurement/comparisons/{payload['id']}/report")
        self.assertEqual(report_response.status_code, 200)
        report = load_workbook(io.BytesIO(report_response.content), read_only=True)
        self.assertEqual(report["报价对比报告"]["A2"].value, "涨价")
        report.close()


if __name__ == "__main__":
    unittest.main()

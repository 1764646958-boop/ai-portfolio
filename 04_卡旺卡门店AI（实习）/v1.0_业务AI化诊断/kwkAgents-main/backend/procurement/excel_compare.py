from __future__ import annotations

from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SHEET_NAME = "装饰直接费"
PRICE_COLUMN = 6  # F
MONEY_QUANTUM = Decimal("0.01")
PERCENT_QUANTUM = Decimal("0.01")
EXCLUDED_ITEM_NAMES = {"小计", "合计", "总计"}
STATUS_LABELS = {
    "increase": "涨价",
    "decrease": "降价",
    "unchanged": "未变化",
    "pending_baseline": "待核定价",
    "new_item": "新增采购项",
    "missing_quote": "报价缺失",
}


class WorkbookValidationError(ValueError):
    """报价工作簿不符合标准清单结构。"""


def _text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\r", "\n").split())


def _money(value: Any) -> Decimal | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped or stripped.startswith("="):
            return None
        stripped = stripped.replace(",", "").replace("￥", "").replace("¥", "")
    else:
        stripped = str(value)
    try:
        amount = Decimal(stripped)
    except (InvalidOperation, ValueError) as exc:
        raise WorkbookValidationError(f"价格 {value!r} 不是有效数字") from exc
    if not amount.is_finite() or amount < 0:
        raise WorkbookValidationError(f"价格 {value!r} 不是有效的非负金额")
    return amount.quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP)


def _money_text(value: Decimal | None) -> str | None:
    return None if value is None else format(value, ".2f")


def _find_header_row(ws) -> int:
    for row_idx in range(1, min(ws.max_row, 40) + 1):
        if (
            _text(ws.cell(row_idx, 1).value) == "序号"
            and _text(ws.cell(row_idx, 2).value) == "项目名称"
            and _text(ws.cell(row_idx, PRICE_COLUMN).value) == "价格"
        ):
            return row_idx
    raise WorkbookValidationError("“装饰直接费”中未找到 A列序号、B列项目名称、F列价格的表头")


def _is_detail_row(ws, row_idx: int) -> bool:
    name = _text(ws.cell(row_idx, 2).value)
    unit = _text(ws.cell(row_idx, 4).value)
    if not name or not unit:
        return False
    compact_name = name.replace(" ", "")
    return compact_name not in EXCLUDED_ITEM_NAMES and not any(
        compact_name.endswith(label) for label in EXCLUDED_ITEM_NAMES
    )


def _load_target_sheet(file_path: str | Path):
    try:
        workbook = load_workbook(
            filename=str(file_path),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise WorkbookValidationError(f"无法读取 Excel 文件：{exc}") from exc
    if SHEET_NAME not in workbook.sheetnames:
        workbook.close()
        raise WorkbookValidationError(f"Excel 中缺少“{SHEET_NAME}”工作表")
    return workbook, workbook[SHEET_NAME]


def create_standard_snapshot(file_path: str | Path) -> dict[str, Any]:
    workbook, ws = _load_target_sheet(file_path)
    try:
        header_row = _find_header_row(ws)
        items: list[dict[str, Any]] = []
        for row_idx in range(header_row + 1, ws.max_row + 1):
            if not _is_detail_row(ws, row_idx):
                continue
            price = _money(ws.cell(row_idx, PRICE_COLUMN).value)
            items.append(
                {
                    "row": row_idx,
                    "sequence": _text(ws.cell(row_idx, 1).value),
                    "name": _text(ws.cell(row_idx, 2).value),
                    "description": _text(ws.cell(row_idx, 3).value),
                    "unit": _text(ws.cell(row_idx, 4).value),
                    "baseline_price": _money_text(price),
                }
            )
        if not items:
            raise WorkbookValidationError("“装饰直接费”中未找到可比较的明细采购项")
        return {
            "format_version": 1,
            "sheet_name": SHEET_NAME,
            "price_column": "F",
            "header_row": header_row,
            "max_row": ws.max_row,
            "items": items,
        }
    finally:
        workbook.close()


def _validate_identity(ws, standard_item: dict[str, Any]) -> None:
    row_idx = int(standard_item["row"])
    actual = {
        "name": _text(ws.cell(row_idx, 2).value),
        "description": _text(ws.cell(row_idx, 3).value),
        "unit": _text(ws.cell(row_idx, 4).value),
    }
    expected = {key: _text(standard_item.get(key)) for key in actual}
    if actual != expected:
        raise WorkbookValidationError(
            f"第 {row_idx} 行项目结构被修改：标准项目“{expected['name']}”，报价文件为“{actual['name']}”"
        )


def _comparison_item(
    standard_item: dict[str, Any],
    quoted_price: Decimal | None,
) -> dict[str, Any]:
    baseline = _money(standard_item.get("baseline_price"))
    status = "unchanged"
    difference: Decimal | None = None
    percentage: Decimal | None = None

    if baseline is None or baseline == 0:
        if quoted_price is not None and quoted_price > 0:
            status = "pending_baseline"
    elif quoted_price is None:
        status = "missing_quote"
    else:
        difference = (quoted_price - baseline).quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP)
        if difference > 0:
            status = "increase"
        elif difference < 0:
            status = "decrease"
        if difference != 0:
            percentage = ((difference / baseline) * Decimal("100")).quantize(
                PERCENT_QUANTUM,
                rounding=ROUND_HALF_UP,
            )

    return {
        "row": int(standard_item["row"]),
        "sequence": _text(standard_item.get("sequence")),
        "name": _text(standard_item.get("name")),
        "description": _text(standard_item.get("description")),
        "unit": _text(standard_item.get("unit")),
        "baseline_price": _money_text(baseline),
        "quoted_price": _money_text(quoted_price),
        "difference": _money_text(difference),
        "percentage": None if percentage is None else format(percentage, ".2f"),
        "status": status,
        "status_label": STATUS_LABELS[status],
    }


def compare_quote_workbook(file_path: str | Path, standard_snapshot: dict[str, Any]) -> dict[str, Any]:
    if int(standard_snapshot.get("format_version", 0)) != 1:
        raise WorkbookValidationError("不支持的标准清单快照版本")

    workbook, ws = _load_target_sheet(file_path)
    try:
        header_row = _find_header_row(ws)
        if header_row != int(standard_snapshot.get("header_row", 0)):
            raise WorkbookValidationError("报价文件的表头位置与标准清单不一致")

        results: list[dict[str, Any]] = []
        for standard_item in standard_snapshot.get("items", []):
            _validate_identity(ws, standard_item)
            row_idx = int(standard_item["row"])
            quoted_price = _money(ws.cell(row_idx, PRICE_COLUMN).value)
            results.append(_comparison_item(standard_item, quoted_price))

        standard_max_row = int(standard_snapshot.get("max_row", 0))
        for row_idx in range(standard_max_row + 1, ws.max_row + 1):
            if not _is_detail_row(ws, row_idx):
                continue
            quoted_price = _money(ws.cell(row_idx, PRICE_COLUMN).value)
            results.append(
                {
                    "row": row_idx,
                    "sequence": _text(ws.cell(row_idx, 1).value),
                    "name": _text(ws.cell(row_idx, 2).value),
                    "description": _text(ws.cell(row_idx, 3).value),
                    "unit": _text(ws.cell(row_idx, 4).value),
                    "baseline_price": None,
                    "quoted_price": _money_text(quoted_price),
                    "difference": None,
                    "percentage": None,
                    "status": "new_item",
                    "status_label": STATUS_LABELS["new_item"],
                }
            )

        summary = {status: 0 for status in STATUS_LABELS}
        for item in results:
            summary[item["status"]] += 1
        summary["total"] = len(results)
        summary["changed"] = summary["increase"] + summary["decrease"]
        return {"format_version": 1, "summary": summary, "items": results}
    finally:
        workbook.close()


def write_comparison_report(result: dict[str, Any], output_path: str | Path) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "报价对比报告"
    headers = ["状态", "原表行号", "序号", "项目名称", "项目特征描述", "单位", "甲方核定价", "供应商报价", "差额", "涨跌幅"]
    ws.append(headers)

    fills = {
        "increase": PatternFill("solid", fgColor="FCE8E6"),
        "decrease": PatternFill("solid", fgColor="E6F4EA"),
        "pending_baseline": PatternFill("solid", fgColor="FFF4CE"),
        "new_item": PatternFill("solid", fgColor="E8F0FE"),
        "missing_quote": PatternFill("solid", fgColor="F3E8FD"),
    }
    for item in result.get("items", []):
        ws.append(
            [
                item.get("status_label"),
                item.get("row"),
                item.get("sequence"),
                item.get("name"),
                item.get("description"),
                item.get("unit"),
                item.get("baseline_price"),
                item.get("quoted_price"),
                item.get("difference"),
                None if item.get("percentage") is None else f"{item['percentage']}%",
            ]
        )
        fill = fills.get(item.get("status"))
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill

    header_fill = PatternFill("solid", fgColor="3D352E")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [14, 10, 10, 34, 54, 10, 14, 14, 14, 12]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)

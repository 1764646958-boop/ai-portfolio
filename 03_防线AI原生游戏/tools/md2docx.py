# -*- coding: utf-8 -*-
"""md → docx 转换工具（团队读 Word，不读 md）。

用法（在项目根目录）：
    python tools/md2docx.py docs/xxx.md 02_yyy.md ...

每个输入的 .md 会在同目录生成同名 .docx。
"""
import io
import os
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

CN_FONT = "微软雅黑"
MONO_FONT = "Consolas"


def set_font(run, size=None, bold=None, italic=None, color=None, mono=False):
    fname = MONO_FONT if mono else CN_FONT
    run.font.name = fname
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), fname)
    rFonts.set(qn('w:hAnsi'), fname)
    rFonts.set(qn('w:eastAsia'), CN_FONT)


INLINE_RE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')


def add_inline(paragraph, text, size=None, bold=False, color=None):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, bold=True, color=color)
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            set_font(run, size=size, bold=bold, color=color, mono=True)
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size, bold=bold, color=color)


def is_table_sep(line):
    t = line.strip()
    return bool(re.match(r'^\|[\s:\-|]+\|$', t)) and '-' in t


def parse_cells(line):
    t = line.strip()
    if t.startswith('|'):
        t = t[1:]
    if t.endswith('|'):
        t = t[:-1]
    return [c.strip() for c in t.split('|')]


def convert(md_rel):
    md_path = md_rel if os.path.isabs(md_rel) else os.path.join(BASE, md_rel)
    docx_path = md_path[:-3] + '.docx'
    with io.open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:eastAsia'), CN_FONT)

    i, n = 0, len(lines)
    while i < n:
        stripped = lines[i].strip()

        if stripped.startswith('```'):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(6)
            for idx, cl in enumerate(code_lines):
                if idx > 0:
                    p.add_run().add_break()
                set_font(p.add_run(cl), size=9.5, mono=True)
            continue

        if stripped.startswith('|') and i + 1 < n and is_table_sep(lines[i + 1]):
            header = parse_cells(lines[i])
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                rows.append(parse_cells(lines[i]))
                i += 1
            ncol = len(header)
            table = doc.add_table(rows=1, cols=ncol)
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            for c in range(ncol):
                add_inline(hdr[c].paragraphs[0], header[c] if c < len(header) else '', size=10, bold=True)
            for row in rows:
                cells = table.add_row().cells
                for c in range(ncol):
                    add_inline(cells[c].paragraphs[0], row[c] if c < len(row) else '', size=10)
            continue

        if stripped.startswith('### '):
            p = doc.add_paragraph()
            set_font(p.add_run(stripped[4:]), size=13, bold=True, color=(0x1F, 0x3B, 0x63))
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            set_font(p.add_run(stripped[3:]), size=15, bold=True, color=(0x1F, 0x3B, 0x63))
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
        elif stripped.startswith('# '):
            p = doc.add_paragraph()
            set_font(p.add_run(stripped[2:]), size=18, bold=True, color=(0x1F, 0x3B, 0x63))
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
        elif stripped == '---':
            pass
        elif stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_inline(p, stripped[2:], size=10.5)
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            add_inline(p, stripped[2:], size=10.5, color=(0x59, 0x59, 0x59))
        elif stripped == '':
            pass
        else:
            p = doc.add_paragraph()
            add_inline(p, stripped, size=10.5)
            p.paragraph_format.space_after = Pt(4)
        i += 1

    doc.save(docx_path)
    return docx_path


def main(argv):
    if not argv:
        print("用法: python tools/md2docx.py <md文件相对路径> ...")
        return
    for md in argv:
        try:
            print("OK  ->", convert(md))
        except Exception as e:  # noqa: BLE001
            print("FAIL->", md, repr(e))


if __name__ == "__main__":
    main(sys.argv[1:])
